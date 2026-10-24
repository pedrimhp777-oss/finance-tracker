import os
from io import BytesIO
from fastapi import FastAPI, Depends, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from scalar_fastapi import get_scalar_api_reference
from dotenv import load_dotenv
from sqlalchemy import extract, text
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.database import engine, Base, get_db
from app import models
from app.routes import transactions, auth
from app.security import get_current_user

load_dotenv()

Base.metadata.create_all(bind=engine)

app = FastAPI(title="Finance Tracker API", docs_url=None, redoc_url=None)

# Declaração explícita das origens para permitir envio de tokens/credenciais
origins = [
    "http://localhost:5173",
    "http://127.0.0.1:5173",
    os.getenv("FRONTEND_URL", "")
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=[o for o in origins if o],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.include_router(auth.router)
app.include_router(transactions.router)


@app.get("/", include_in_schema=False)
def home():
    return {"message": "API de Finanças rodando com sucesso!"}


@app.get("/health", include_in_schema=False)
def health(db: Session = Depends(get_db)):
    db.execute(text("SELECT 1"))
    return {"status": "ok"}

@app.get("/docs", include_in_schema=False)
def scalar_html():
    return get_scalar_api_reference(
        openapi_url=app.openapi_url,
        title=app.title,
    )

@app.get("/transactions/export/excel")
def export_excel(
    month: int | None = Query(None, ge=1, le=12),
    year: int | None = Query(None, ge=1, le=9999),
    category: str | None = Query(None, min_length=1, max_length=100),
    search: str | None = Query(None, max_length=240),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    query = db.query(models.Transaction).filter(models.Transaction.user_id == current_user.id)
    if month is not None:
        query = query.filter(extract("month", models.Transaction.date) == month)
    if year is not None:
        query = query.filter(extract("year", models.Transaction.date) == year)
    if category:
        query = query.filter(models.Transaction.category == category)
    if search:
        query = query.filter(models.Transaction.description.ilike(f"%{search.strip()}%"))
    user_transactions = query.order_by(models.Transaction.date.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório Financeiro"

    headers = ["Descrição", "Categoria", "Tipo", "Valor (R$)", "Data"]
    ws.append(headers)

    header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    for t in user_transactions:
        ws.append([
            t.description,
            t.category,
            "Entrada" if t.type == "income" else "Saída",
            float(t.amount),
            t.date.strftime("%d/%m/%Y") if hasattr(t.date, 'strftime') else str(t.date)
        ])

    currency_format = 'R$ #,##0.00'
    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )

    for row in ws.iter_rows(min_row=2, max_row=len(user_transactions) + 1, min_col=1, max_col=5):
        for cell in row:
            cell.border = thin_border
            cell.font = Font(name="Segoe UI", size=10)
        
        row[2].alignment = center_align
        row[4].alignment = center_align
        row[3].number_format = currency_format

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 6, 14)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = "relatorio_financeiro.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

